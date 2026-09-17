"""Parses uploaded files into LangChain Documents.

Each PDF page becomes its own Document (so page numbers survive into chunk
metadata for citations); DOCX/TXT don't have a reliable page concept, so the
whole file becomes a single Document.
"""

from pathlib import Path

import docx
from langchain_core.documents import Document
from pypdf import PdfReader

from .config import ALLOWED_EXTENSIONS
from .errors import DocumentParsingError, EmptyDocumentError, UnsupportedFileTypeError


def _load_pdf(path: Path) -> list[Document]:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentParsingError(f"Could not read PDF: {exc}") from exc

    docs = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(page_content=text, metadata={"page": page_number}))
    return docs


def _load_docx(path: Path) -> list[Document]:
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocumentParsingError(f"Could not read DOCX: {exc}") from exc

    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    return [Document(page_content=text)] if text.strip() else []


def _load_txt(path: Path) -> list[Document]:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except OSError as exc:
        raise DocumentParsingError(f"Could not read text file: {exc}") from exc
    return [Document(page_content=text)] if text.strip() else []


_LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_txt}


def parse_file(path: Path, filename: str) -> list[Document]:
    """Parse a saved upload into one Document per page/section.

    Raises UnsupportedFileTypeError, DocumentParsingError, or
    EmptyDocumentError — callers turn these into 4xx responses.
    """
    ext = Path(filename).suffix.lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext or filename}'. "
            f"Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    docs = loader(path)
    if not docs:
        raise EmptyDocumentError(f"No extractable text found in '{filename}'.")
    return docs
